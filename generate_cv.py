from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import weasyprint
import os

# ── CV DATA ────────────────────────────────────────────────────────────────────

NAME = "IBRAHEEM MUSTAFA"
TITLE = "Information Security Engineer – GRC"
CONTACT = [
    "ibraheem.mustafa00@gmail.com",
    "+92-3018442580",
    "Lahore, Pakistan",
    "LinkedIn: Ibraheem Mustafa",
]

SUMMARY = (
    "GRC-focused Information Security Engineer with 2+ years of experience delivering "
    "ISO 27001:2022 audit engagements, ISMS implementations, and compliance assessments "
    "across local and international organizations. Conducted end-to-end internal and "
    "external audits for 5 organizations including Kualitatem, Virtual Force, and "
    "international clients, driving ISMS certification and compliance outcomes. Certified "
    "ISO 27001:2022 Lead Auditor (CQI|IRCA & Mastermind Assurance) with hands-on "
    "expertise in SAMA CSF, SOC 2, GDPR, Vanta, and Sahl."
)

SKILLS = (
    "ISO 27001:2022  |  Internal Audit  |  External Audit  |  ISMS Implementation  |  "
    "SAMA CSF  |  SOC 2  |  GDPR  |  Risk Assessment  |  Risk Treatment Plan  |  "
    "Gap Analysis  |  Statement of Applicability (SoA)  |  Non-Conformity Reporting  |  "
    "Audit Planning & Execution  |  Compliance Reporting  |  Vanta  |  Sahl  |  "
    "Top Management Presentations"
)

EXPERIENCE = [
    {
        "company": "Kualitatem Inc.",
        "title": "Information Security Engineer – Information Security Department",
        "dates": "December 2023 – Present",
        "location": "Lahore, Pakistan",
        "bullets": [
            "Strengthened Kualitatem's ISO 27001:2022 compliance posture, as measured by a successful external certification audit with zero major non-conformities, by leading end-to-end internal audit activities covering planning, evidence collection, controls assessment, and non-conformity reporting to top management.",
            "Maintained ISO 27001:2022 certification for Virtual Force, as measured by a passed surveillance audit, by executing a full audit lifecycle including scope review, controls re-evaluation, and verification of previously identified non-conformity closures.",
            "Achieved ISO 27001:2022 certification readiness for 3 clients (2 local, 1 international), as measured by successful audit outcomes across all engagements, by delivering structured end-to-end internal and external audit programs, identifying control gaps, and presenting prioritized findings to client management.",
            "Achieved ISO 27001:2022 certification for a client from ground zero, as measured by full certification attainment, by leading end-to-end ISMS implementation including risk assessment, SoA development, control design, and pre-certification audit preparation.",
            "Renewed SOC 2 certification for Kualitatem, as measured by continued compliance with all Trust Service Criteria, by coordinating cross-functional evidence collection via Vanta, liaising with external auditors, and remediating identified control gaps ahead of assessment.",
            "Enhanced client compliance posture against SAMA CSF, as measured by detailed gap assessment reports delivered to senior management, by conducting structured stakeholder interviews, analyzing control deficiencies via Sahl, and producing prioritized remediation roadmaps.",
            "Supported organization-wide GDPR compliance program, as measured by documented data processing assessments, by collaborating with the Internal Audit team to review data handling practices and implement corrective control measures.",
        ],
    }
]

EDUCATION = [
    ("Masters in Cyber Security – CGPA 3.57", "Bahcesehir University, Turkiye", "2023"),
    ("Bachelors in Computer Science – CGPA 3.40", "Bahauddin Zakariya University, Pakistan", "2018"),
]

CERTIFICATIONS = [
    ("ISO 27001:2022 Lead Auditor", "CQI|IRCA"),
    ("ISO/IEC 27001:2022 Lead Auditor", "Mastermind Assurance"),
    ("ISO 27001:2022 INFOSEC Associate", "SkillFront"),
    ("eJPT – Junior Penetration Tester", "INE eLearnSecurity"),
]

LANGUAGES = [
    ("English", "Full Professional Proficiency"),
    ("Urdu", "Native / Bilingual Proficiency"),
    ("Turkish", "Elementary Proficiency"),
]

# ── NexGen GRC Color Palette ──────────────────────────────────────────────────
# Dark Navy   #1E3D7A  (Forward Looking)
# Teal        #28AFA0  (Adequacy)
# Deep Navy   #1A2D60  (Relevancy)
# Purple      #6B35A0  (Timeliness)
# Amber       #F5A520  (Actionability)
# Crimson     #C0232B  (Center / accent)

C_DARK_NAVY  = RGBColor(0x1E, 0x3D, 0x7A)
C_TEAL       = RGBColor(0x28, 0xAF, 0xA0)
C_DEEP_NAVY  = RGBColor(0x1A, 0x2D, 0x60)
C_PURPLE     = RGBColor(0x6B, 0x35, 0xA0)
C_AMBER      = RGBColor(0xF5, 0xA5, 0x20)
C_CRIMSON    = RGBColor(0xC0, 0x23, 0x2B)
C_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT_GRAY = RGBColor(0xF2, 0xF4, 0xF8)

HEX_DARK_NAVY = "1E3D7A"
HEX_TEAL      = "28AFA0"
HEX_DEEP_NAVY = "1A2D60"
HEX_PURPLE    = "6B35A0"
HEX_AMBER     = "F5A520"
HEX_CRIMSON   = "C0232B"
HEX_LIGHT_BG  = "F2F4F8"

# ── DOCX GENERATION ────────────────────────────────────────────────────────────

def set_cell_bg(cell, color_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def add_colored_hrule(doc, color_hex="28AFA0"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)

def section_heading(doc, text, rule_color=HEX_TEAL):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = C_DARK_NAVY
    add_colored_hrule(doc, rule_color)

def build_docx(path):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)

    # ── Header band ──
    # Name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(NAME)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = C_DEEP_NAVY

    # Title with teal color
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(4)
    run2 = p2.add_run(TITLE)
    run2.font.size = Pt(12)
    run2.font.color.rgb = C_TEAL
    run2.bold = True

    # Contact line
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(6)
    run3 = p3.add_run("  |  ".join(CONTACT))
    run3.font.size = Pt(9)
    run3.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # Thick double rule: amber then teal
    add_colored_hrule(doc, HEX_AMBER)
    add_colored_hrule(doc, HEX_TEAL)

    # ── Summary ──
    section_heading(doc, "Professional Summary", HEX_TEAL)
    p = doc.add_paragraph(SUMMARY)
    p.paragraph_format.space_after = Pt(4)

    # ── Skills (pill-style row with purple heading) ──
    section_heading(doc, "Core Skills", HEX_PURPLE)
    p = doc.add_paragraph(SKILLS)
    p.paragraph_format.space_after = Pt(4)

    # ── Experience ──
    section_heading(doc, "Work Experience", HEX_DARK_NAVY)
    for job in EXPERIENCE:
        # Company name in dark navy
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(job["company"])
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = C_DARK_NAVY

        # Job title in purple
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(job["title"])
        r2.bold = True
        r2.font.size = Pt(10)
        r2.font.color.rgb = C_PURPLE

        # Dates in amber
        p3 = doc.add_paragraph()
        p3.paragraph_format.space_before = Pt(0)
        p3.paragraph_format.space_after = Pt(4)
        r3 = p3.add_run(f"{job['dates']}  |  {job['location']}")
        r3.italic = True
        r3.font.size = Pt(9)
        r3.font.color.rgb = C_AMBER

        for bullet in job["bullets"]:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(bullet)
            run.font.size = Pt(10)

    # ── Education (teal heading) ──
    section_heading(doc, "Education", HEX_TEAL)
    for degree, uni, year in EDUCATION:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(degree)
        r.bold = True
        r.font.color.rgb = C_DEEP_NAVY
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(f"{uni}  |  {year}")
        r2.font.size = Pt(9)
        r2.italic = True
        r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # ── Certifications table (dark navy header) ──
    section_heading(doc, "Certifications", HEX_CRIMSON)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = "Certification"
    hdr[1].text = "Issuing Body"
    for cell in hdr:
        set_cell_bg(cell, HEX_DEEP_NAVY)
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.color.rgb = C_WHITE
                run.font.size = Pt(10)
    for i, (cert, issuer) in enumerate(CERTIFICATIONS):
        row = table.add_row().cells
        row[0].text = cert
        row[1].text = issuer
        bg = HEX_LIGHT_BG if i % 2 == 0 else "FFFFFF"
        set_cell_bg(row[0], bg)
        set_cell_bg(row[1], bg)
        for cell in row:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)

    doc.add_paragraph()

    # ── Languages ──
    section_heading(doc, "Languages", HEX_AMBER)
    for lang, prof in LANGUAGES:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"{lang}: ")
        r1.bold = True
        r1.font.color.rgb = C_DARK_NAVY
        p.add_run(prof)

    doc.save(path)
    print(f"DOCX saved: {path}")


# ── HTML / PDF GENERATION ──────────────────────────────────────────────────────

HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8"/>
<style>
  /* NexGen GRC Palette
     Dark Navy  #1E3D7A  Forward Looking
     Teal       #28AFA0  Adequacy
     Deep Navy  #1A2D60  Relevancy
     Purple     #6B35A0  Timeliness
     Amber      #F5A520  Actionability
     Crimson    #C0232B  Center
  */
  @page {{ margin: 1.5cm 2cm; }}
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: Calibri, Arial, sans-serif; font-size: 10pt; color: #222; line-height: 1.45; }}

  /* ── Header ── */
  .header {{
    border-top: 4px solid #F5A520;
    border-bottom: 3px solid #28AFA0;
    padding: 10px 0 8px 0;
    text-align: center;
    margin-bottom: 12px;
  }}
  .header h1 {{
    font-size: 22pt;
    color: #1A2D60;
    letter-spacing: 2px;
    margin-bottom: 3px;
    text-transform: uppercase;
  }}
  .header .title {{
    font-size: 12pt;
    color: #28AFA0;
    font-weight: bold;
    margin-bottom: 5px;
  }}
  .header .contact {{
    font-size: 9pt;
    color: #555;
  }}

  /* ── Section headings ── */
  .section-title {{
    font-size: 11pt;
    font-weight: bold;
    color: #1E3D7A;
    text-transform: uppercase;
    letter-spacing: 0.8px;
    margin-top: 14px;
    margin-bottom: 5px;
    padding-bottom: 3px;
    border-bottom: 2px solid #28AFA0;
  }}
  .section-title.purple  {{ border-color: #6B35A0; }}
  .section-title.navy    {{ border-color: #1A2D60; }}
  .section-title.crimson {{ border-color: #C0232B; }}
  .section-title.amber   {{ border-color: #F5A520; }}

  /* ── Summary ── */
  .summary {{ margin-bottom: 4px; }}

  /* ── Skills ── */
  .skills {{
    background: #F2F4F8;
    border-left: 4px solid #6B35A0;
    padding: 6px 10px;
    margin-bottom: 4px;
    font-size: 9.5pt;
    color: #333;
  }}

  /* ── Experience ── */
  .job {{ margin-bottom: 10px; }}
  .job-company {{
    font-size: 11pt;
    font-weight: bold;
    color: #1E3D7A;
    margin-top: 8px;
    margin-bottom: 1px;
  }}
  .job-title {{
    font-weight: bold;
    font-size: 10pt;
    color: #6B35A0;
    margin-bottom: 1px;
  }}
  .job-meta {{
    font-size: 9pt;
    color: #F5A520;
    font-style: italic;
    font-weight: bold;
    margin-bottom: 5px;
  }}
  ul {{
    margin-left: 18px;
    margin-bottom: 4px;
  }}
  li {{
    margin-bottom: 3px;
    padding-left: 2px;
  }}
  li::marker {{ color: #28AFA0; }}

  /* ── Education ── */
  .edu-entry {{ margin-bottom: 7px; }}
  .edu-degree {{ font-weight: bold; color: #1A2D60; }}
  .edu-uni {{ font-size: 9pt; color: #666; font-style: italic; }}

  /* ── Certifications table ── */
  table {{
    width: 100%;
    border-collapse: collapse;
    margin-top: 5px;
    font-size: 10pt;
  }}
  th {{
    background: #1A2D60;
    color: #fff;
    font-weight: bold;
    padding: 6px 10px;
    text-align: left;
  }}
  td {{
    border: 1px solid #ddd;
    padding: 5px 10px;
  }}
  tr:nth-child(even) td {{ background: #F2F4F8; }}
  tr:nth-child(odd)  td {{ background: #fff; }}

  /* ── Languages ── */
  .lang-row {{ margin-bottom: 3px; }}
  .lang-name {{ font-weight: bold; color: #1E3D7A; }}

  /* ── Footer stripe ── */
  .footer-stripe {{
    margin-top: 16px;
    height: 4px;
    background: linear-gradient(to right,
      #1E3D7A 0%, #1E3D7A 20%,
      #28AFA0 20%, #28AFA0 40%,
      #6B35A0 40%, #6B35A0 60%,
      #F5A520 60%, #F5A520 80%,
      #C0232B 80%, #C0232B 100%);
  }}
</style>
</head>
<body>

<div class="header">
  <h1>{name}</h1>
  <div class="title">{title}</div>
  <div class="contact">{contact}</div>
</div>

<div class="section-title">Professional Summary</div>
<p class="summary">{summary}</p>

<div class="section-title purple">Core Skills</div>
<div class="skills">{skills}</div>

<div class="section-title navy">Work Experience</div>
{experience}

<div class="section-title">Education</div>
{education}

<div class="section-title crimson">Certifications</div>
<table>
  <tr><th>Certification</th><th>Issuing Body</th></tr>
  {cert_rows}
</table>

<div class="section-title amber">Languages</div>
{languages}

<div class="footer-stripe"></div>

</body>
</html>"""


def build_pdf(html_path, pdf_path):
    weasyprint.HTML(filename=html_path).write_pdf(pdf_path)
    print(f"PDF saved: {pdf_path}")


def build_html():
    contact_str = "  |  ".join(CONTACT)

    exp_html = ""
    for job in EXPERIENCE:
        bullets = "".join(f"<li>{b}</li>" for b in job["bullets"])
        exp_html += f"""
        <div class="job">
          <div class="job-company">{job['company']}</div>
          <div class="job-title">{job['title']}</div>
          <div class="job-meta">{job['dates']}  |  {job['location']}</div>
          <ul>{bullets}</ul>
        </div>
        """

    edu_html = ""
    for degree, uni, year in EDUCATION:
        edu_html += f"""
        <div class="edu-entry">
          <div class="edu-degree">{degree}</div>
          <div class="edu-uni">{uni}  |  {year}</div>
        </div>
        """

    cert_rows = "".join(f"<tr><td>{c}</td><td>{i}</td></tr>" for c, i in CERTIFICATIONS)

    lang_html = "".join(
        f'<div class="lang-row"><span class="lang-name">{l}:</span> {p}</div>'
        for l, p in LANGUAGES
    )

    return HTML_TEMPLATE.format(
        name=NAME,
        title=TITLE,
        contact=contact_str,
        summary=SUMMARY,
        skills=SKILLS,
        experience=exp_html,
        education=edu_html,
        cert_rows=cert_rows,
        languages=lang_html,
    )


if __name__ == "__main__":
    out_dir = "/home/user/CV-Creator---Claude-AI/output"
    os.makedirs(out_dir, exist_ok=True)

    docx_path = f"{out_dir}/Ibraheem_Mustafa_GRC_CV.docx"
    html_path = f"{out_dir}/Ibraheem_Mustafa_GRC_CV.html"
    pdf_path  = f"{out_dir}/Ibraheem_Mustafa_GRC_CV.pdf"

    build_docx(docx_path)

    html = build_html()
    with open(html_path, "w") as f:
        f.write(html)
    print(f"HTML saved: {html_path}")

    build_pdf(html_path, pdf_path)
    print("Done.")
